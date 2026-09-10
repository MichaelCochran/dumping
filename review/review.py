import csv
import glob
import hashlib
import json
import os
import re
import shutil
import sys
import xml.etree.ElementTree as ET
from enum import Enum
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.borders import Border, Side

pd.options.mode.chained_assignment = None

GREEN = '\033[92m'
YELLOW = '\033[93m'
RED = '\033[91m'
WHITE = '\033[97m'
ENDCOLOR = '\033[0m'

os.system('')

integrationPath = 'integration_reports'
featurePath = 'feature_reports'
commonPath = 'common_report'

MANIFEST_NAME = 'manifest.csv'

thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))


def sizeColumns(**sizeInfo):
    for x in range(sizeInfo['numCols']):
        currCol = chr(ord('a') + x)
        sizeInfo['sheetName'].column_dimensions[currCol].width = sizeInfo[currCol]


def alignXlsCells(sheetName, colList, horizontal, vertical, wrap, border):
    for col in colList:
        for cell in sheetName[col]:
            cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrapText=wrap)
            if border == True:
                cell.border = thin_border


def createDF_func(fileDirectory, branchType, hasComments):
    global branch
    if branchType == 'integration':
        for f1 in fileDirectory:
            if str(repoTag1) in f1.lower():
                global integrationReportPath
                integrationReportPath = integrationPath + '/' + f1
                if hasComments:
                    integrationDF = pd.read_csv(integrationReportPath, usecols=['Category', 'Full Filename', 'Line Number', 'CWE', 'Tagged', 'Criticality', 'Primary Location', 'Instance ID', 'Comment'])
                    integrationDF = integrationDF.loc[:, ['Category', 'CWE', 'Line Number', 'Primary Location', 'Tagged', 'Criticality', 'Comment', 'Full Filename', 'Instance ID']]
                else:
                    integrationDF = pd.read_csv(integrationReportPath, usecols=['Category', 'Full Filename', 'Line Number', 'CWE', 'Tagged', 'Criticality', 'Primary Location', 'Instance ID'])
                    integrationDF = integrationDF.loc[:, ['Category', 'CWE', 'Line Number', 'Primary Location', 'Tagged', 'Criticality', 'Full Filename', 'Instance ID']]

                integrationDF = integrationDF.sort_values(by=['Full Filename', 'Line Number'])
                integrationDF['CWE'] = integrationDF['CWE'].str.split(',').str[0]

                integrationDF = integrationDF.reset_index(drop=True)
                integrationDF = integrationDF.drop_duplicates(subset=['CWE', 'Line Number', 'Primary Location'])
                return integrationDF
        if not 'integrationReportPath' in globals():
            print(f"{RED} No integration file in {repoTag1} found. Please download and export. Try again. {ENDCOLOR}")
            return

    elif branchType == 'feature':
        for f2 in featureFiles:
            if repoTag1 in f2:
                if not (re.match(r'\S{3,}_\d{4}-\d{2}\d{2}\.csv$', f2)):
                    print(f"{RED}Format of source file is incorrect.{ENDCOLOR}")
                    exit()

                featureBranch = f2.replace('feature_reports\\', '')
                featureBranch = featureBranch[:-15]
                branch = re.sub(r'^[^_]*_', '', featureBranch)
                for ch in ('\\', '/', '_'):
                    branch = branch.replace(ch, '-')
                try:
                    input(f"\nGenerating report for branch {GREEN}{featureBranch}{ENDCOLOR}.\nPress [Enter] to continue or [Ctrl + C] to cancel: ")
                except KeyboardInterrupt:
                    exit()

                global featReportPath
                featReportPath = f2
                if hasComments:
                    featDF = pd.read_csv(f2, usecols=['Category', 'Full Filename', 'Line Number', 'Criticality', 'CWE', 'Tagged', 'Primary Location', 'Instance ID', 'Comment'])
                else:
                    featDF = pd.read_csv(f2, usecols=['Category', 'Full Filename', 'Line Number', 'Criticality', 'CWE', 'Tagged', 'Primary Location', 'Instance ID'])

                featDF = featDF.drop_duplicates(subset=['Category', 'CWE', 'Primary Location', 'Line Number', 'Criticality'])
                featDF = featDF.rename(columns={'Criticality': 'Fortify Criticality'})

                global workingFeatFile
                workingFeatFile = f2[16:]
                featDF['CWE'] = featDF['CWE'].str.split(',').str[0]

                if repoTag1 != 'common':
                    featDF = featDF[~featDF['Full Filename'].str.contains('common_fsw')]
                    featDF = featDF[~featDF['Full Filename'].str.contains('orca')]

                    commonFiles = os.listdir(commonPath)
                    for f3 in commonFiles:
                        commonReportPath = commonPath + '/' + f3

                    commonDF = pd.read_csv(commonReportPath, usecols=['Category', 'Full Filename', 'Line Number', 'CWE', 'Primary Location'])
                    commonDF = commonDF[commonDF['Full Filename'].str.contains('/usr/local/')].reset_index(drop=True)

                    featDF = pd.merge(featDF, commonDF, how='left', indicator='Preexisting StdLib')
                    featDF['Preexisting StdLib'] = np.where(featDF['Preexisting StdLib'] == 'both', True, False)
                    featDF = featDF[~featDF['Preexisting StdLib']].reset_index(drop=True)
                    featDF = featDF.drop_duplicates(subset=['CWE', 'Line Number', 'Primary Location'])

                return featDF
                break
        if not 'featReportPath' in globals():
            print(f"{RED}No Source file in {repoTag1} found. Please download data export from Fortify.\nQuitting...{ENDCOLOR}")
            exit()


def createReport():
    integrationFiles = os.listdir(integrationPath)
    integrationDF = createDF_func(integrationFiles, 'integration', hasComments)
    if integrationDF is None:
        return False

    global featureFiles
    globoPath = featurePath + '/*'
    featureFiles = glob.glob(globoPath)
    featureFiles.sort(key=os.path.getmtime, reverse=True)
    featDF = createDF_func(featureFiles, 'feature', hasComments)

    global reviewDF
    reviewDF = featDF

    id_match = featDF['Instance ID'].isin(integrationDF['Instance ID'])

    integration_key_set = set(integrationDF[['Primary Location', 'Line Number', 'Category']].apply(tuple, axis=1))
    feat_key_series = featDF[['Primary Location', 'Line Number', 'Category']].apply(tuple, axis=1)
    key_match = feat_key_series.isin(integration_key_set)

    tagged = (reviewDF['Tagged'].notna() & reviewDF['Tagged'].astype(str).str.strip().ne(''))

    reviewDF['Preexisting Backlog'] = id_match | key_match | tagged
    if hasComments:
        reviewDF = reviewDF.sort_values(by=['Preexisting Backlog', 'Full Filename', 'Category', 'Comment'],
                                         ascending=[True, True, True, True])
        reviewDF = reviewDF.loc[:, ['Category', 'CWE', 'Line Number', 'Primary Location',
                                     'Tagged', 'Preexisting Backlog', 'Fortify Criticality', 'Comment', 'Full Filename']]
    else:
        reviewDF = reviewDF.sort_values(by=['Preexisting Backlog', 'Full Filename', 'Category'],
                                         ascending=[True, True, True])
        reviewDF = reviewDF.loc[:, ['Category', 'CWE', 'Line Number', 'Primary Location', 'Tagged',
                                     'Preexisting Backlog', 'Fortify Criticality', 'Full Filename']]

    reviewDF = reviewDF.reset_index(drop=True)
    integrationDF = integrationDF.drop(columns=['Instance ID'])

    cweTree = ET.parse('archive/mitre/cwe_data.xml')
    cweRoot = cweTree.getroot()

    DFs = [integrationDF, reviewDF]

    for df in DFs:
        df.insert(2, 'Potential Impact', pd.Series(dtype='str'))
        df.insert(2, 'Description', pd.Series(dtype='str'))
        df['Potential Impact'] = ''
        df['Description'] = ''

        for finding in df.index:
            for weakness in cweRoot.iter('Weakness'):
                if ('CWE ID' + weakness.attrib['ID'] == df['CWE'][finding]):
                    description = weakness.find('Description')
                    df.at[finding, 'Description'] += description.text
                    impacts = weakness.iter('Impact')
                    for impact in impacts:
                        df.at[finding, 'Potential Impact'] += 'Impact: ' + impact.text + '\n'

    global reportName
    reportName = '_cyberReview_' + workingFeatFile[:-4] + '.xlsx'

    with pd.ExcelWriter(reportName, mode='w') as writer:
        reviewDF.to_excel(writer, sheet_name="Review Branch Report")
        integrationDF.to_excel(writer, sheet_name="Integration Branch Report")

    formatReport(reportName)
    return True


def formatReport(report):
    writer = pd.ExcelWriter(reportName, engine='openpyxl', mode='a')
    workbook = writer.book
    sheets = workbook.sheetnames

    for sheet in sheets:
        ws = workbook[sheet]

        for cell in ws['H']:
            if cell.value:
                if 'Suspicious' in cell.value:
                    cell.fill = PatternFill('solid', fgcolor='F08A89')
                elif 'Not Assessed' in cell.value:
                    cell.fill = PatternFill('solid', fgcolor='f5f58c')
                elif 'NA' in cell.value:
                    cell.fill = PatternFill('solid', fgcolor='96faad')
                elif 'Whitelist APPROVED' in cell.value:
                    cell.fill = PatternFill('solid', fgcolor='e6c200')

        if sheet == "Review Branch Report":
            alignXlsCells(sheetName=ws, colList=['A', 'B', 'C', 'F', 'H', 'I', 'J'], horizontal='center', vertical='center', wrap=True, border=True)
            alignXlsCells(sheetName=ws, colList=['D', 'E', 'K', 'L'], horizontal='left', vertical='top', wrap=True, border=True)
            alignXlsCells(sheetName=ws, colList=['G'], horizontal='left', vertical='center', wrap=True, border=True)

            sizeColumns(sheetName=ws, numCols=12, a=5, b=15, c=12, d=50, e=42, f=9, g=22, h=12, i=12, j=10, k=25, l=35)

            for cell in ws['I']:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if not cell.value:
                    cell.fill = PatternFill('solid', fgColor='FFC000')

            for row in ws['A1:J1']:
                for cell in row:
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        elif sheet == "Integration Branch Report":
            alignXlsCells(sheetName=ws, colList=['A', 'B', 'C', 'F', 'H', 'I'], horizontal='center', vertical='center', wrap=True, border=True)
            alignXlsCells(sheetName=ws, colList=['D', 'E', 'J', 'K'], horizontal='left', vertical='top', wrap=True, border=True)
            alignXlsCells(sheetName=ws, colList=['G'], horizontal='left', vertical='center', wrap=True, border=True)

            sizeColumns(sheetName=ws, numCols=12, a=5, b=15, c=12, d=50, e=42, f=9, g=22, h=12, i=12, j=10, k=25, l=35)

            for row in ws['A1:J1']:
                for cell in row:
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        ws.insert_rows(0)
        ws.insert_rows(0)
        numCols = ws.max_column
        numRows = ws.max_row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=numCols)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=numCols)
        ws.merge_cells(start_row=numRows + 1, start_column=1, end_row=numRows + 1, end_column=numCols)
        ws.merge_cells(start_row=numRows + 2, start_column=1, end_row=numRows + 2, end_column=numCols)

        ws['A1'] = "Don't Speak"
        ws['A1'].font = Font(name='Times New Roman', size=10, color='DC143C')
        ws['A1'].alignment = Alignment(horizontal='center')

        footerCell = 'A' + str(numRows + 1)
        ws[footerCell] = "Don't Speak"
        ws[footerCell].font = Font(name='Times New Roman', size=10, color='DC143C')
        ws[footerCell].alignment = Alignment(horizontal='center')

    writer.close()


def get_filename(full_path: str) -> str:
    slash_pos = full_path.rfind('/')
    backslash_pos = full_path.rfind('\\')
    sep_index = max(slash_pos, backslash_pos)
    file_name = full_path[sep_index + 1:] if sep_index != -1 else full_path
    print(f"{file_name} found")
    return file_name


def archiveRptFiles():
    # TODO Use a single temporary folder instead of the archive directory setup below

    archiveParentDir = f'archive/{repoTag1}'
    archiveSubDir = collabReview + '_' + workingFeatFile[:-4] + '/'

    global fullDir
    fullDir = archiveParentDir + archiveSubDir
    old_manifest = readManifest(fullDir)

    try:
        os.makedirs(fullDir)
    except FileExistsError:
        try:
            input(f"{YELLOW}A report folder for this branch already exists. Press [ENTER] to overwrite or [Ctrl +C] to cancel: {ENDCOLOR}")
        except KeyboardInterrupt:
            return False
        try:
            os.makedirs(fullDir, exist_ok=True)
        except PermissionError:
            return False

    srcFiles = [{'src': reportName, 'op': 'move'}, {'src': featReportPath, 'op': 'copy'}, {'src': integrationReportPath, 'op': 'copy'}]

    for entry in srcFiles:
        src = entry['src']
        dst = os.path.join(fullDir, os.path.basename(src))
        try:
            if entry['op'] == 'move':
                shutil.move(src, dst)
            else:
                shutil.copy2(src, dst)
        except PermissionError:
            return False

    return True


def trackFindings():
    newFindings = 0

    for row in reviewDF.index:
        if (reviewDF['Preexisting Backlog'][row] == False):
            newFindings += 1

    if newFindings == 0:
        print(f"{GREEN}\nNo new findings exist! Visual verification is optional.{ENDCOLOR}")
    else:
        print(f"{YELLOW}\nThere are {RED}{newFindings} new findings{YELLOW}. Open the cyberReview file to make sure they are really new.{ENDCOLOR}\nChech out the new findings in this branch {WHITE}{branch}{YELLOW}.{ENDCOLOR}")

    try:
        input(f"\nAfter fixing and mistakes in the cyberReview and saving the file, press [Eter] to continue")
    except KeyboardInterrupt:
        exit()

    trackingDF = pd.read_excel(fullDir + reportName, usecols=['Category', 'CWE', 'Line Number', 'Primary Location', 'Tagged', 'Preexisting Backlog', 'Full Filename'], sheet_name=1, header=2, skipfooter=2)

    trackingDF = trackingDF[~trackingDF['Preexisting Backlog']]
    trackingDF['Notes'] = "Add notes here"
    trackingDF['Review #'] = collabReview

    sumDF = trackingDF.loc[:, ['Category', 'Primary Location', 'Full Filename']]
    sumDF['Count'] = 1
    sumDF = sumDF.groupby(['Category', 'Primary Location', 'Full Filename']).Count.count().reset_index()

    sumDF = sumDF.sort_values(by=['Full Filename', 'Category'])

    finalCount = 0
    for index, row in sumDF.iterrows():
        finalCount + row['Count']
        print(str(row['Count']) + "" + row['Category'] + " Finding(s) in " + row['Primary Location'])

        print(str(finalCount) + " New Findings")


def _hash_file(p: str) -> str:
    if os.path.isfile(p):
        h = hashlib.sha256()
        with open(p, 'rb') as f:
            for chunk in iter(lambda: f.read(8192), b''):
                h.update(chunk)
        return h.hexdigest()
    return ''


def readManifest(report_dir: Path) -> dict:
    manifest_path = os.path.join(report_dir, MANIFEST_NAME)
    data = {}
    if os.path.isfile(manifest_path):
        with open(manifest_path, mode='r', newline='', encoding='utf8') as f:
            for name, h in csv.reader(f):
                data[name] = h
    return data


def configureEnum():
    with open('repoConfig.json', 'r') as f:
        config = json.load(f)

    enumRepo = Enum('enumRepo', config)

    return enumRepo


if __name__ == '__main__':
    enumRepo = configureEnum()
    global repoTag1
    hasComments = False
    repoSet = False
    if len(sys.argv) > 1:
        hasComments = True
        repoSet = True
        repoTag1 = sys.argv[1]
    else:
        comments = input("\nDo your reports have comments? (Y/N) ")
        while True:
            if comments.upper() == 'N':
                break
            elif comments.upper() == 'Y':
                hasComments = True
                break
            else:
                comments = input("Please enter Y or N ")

    while True:
        while True:
            if repoSet == False:
                while True:
                    try:
                        for tag in enumRepo:
                            print(f"{tag.value} - {tag.name.lower()}")

                        userInput = int(input("Please enter a number: "))
                        if userInput < 1 or userInput > len(enumRepo):
                            raise ValueError
                        else:
                            repoTag1 = enumRepo(userInput).name.lower()
                            for ch in ('\\', '/', '_'):
                                repoTag1 = repoTag1.replace(ch, '-')
                            break
                    except ValueError:
                        print(f"{RED}Please enter a valid number!{ENDCOLOR}")

            global collabReview
            collabReview = input("Input the Collaborator Review #: ")
            if createReport():
                break

        if archiveRptFiles():
            break
        else:
            input(f"\n\n{RED}Error writing file. Try closing open Excel files, and check folder permissions. \nPress [Enter] to continue. {ENDCOLOR}")

    trackFindings()

    quit()
